"""Build the complete CreditScope project documentation as a verified Word document."""

from __future__ import annotations

import csv
import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUT = ROOT / "CreditScope_Complete_Project_Documentation.docx"

INK = "17324D"
TEAL = "17766D"
PALE = "EAF4F2"
PALE_BLUE = "EFF5F9"
LIGHT = "F7F9FA"
GRID = "D9D9D9"
MUTED = "5D6B73"
WHITE = "FFFFFF"


def load_csv(name: str) -> list[dict[str, str]]:
    with (DATA / "research" / name).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=95, start=110, bottom=95, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=GRID, size="6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, width: float) -> None:
    cell.width = Inches(width)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(table) -> None:
    if table.rows:
        repeat_header(table.rows[0])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float],
              font_size: float = 8.5, alignments: list[str] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = str(header)
        shade(cell, INK)
        set_col_width(cell, widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_borders(cell)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.runs[0]
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(font_size)
    repeat_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(values):
            cell = cells[col_index]
            cell.text = str(value)
            set_col_width(cell, widths[col_index])
            set_cell_margins(cell)
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                if alignments and alignments[col_index] == "center":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif alignments and alignments[col_index] == "right":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string("1F2933")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = False


def add_number(doc: Document, text: str, level: int = 0) -> None:
    style = "List Number" if level == 0 else "List Number 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(4)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = False


def add_picture(doc: Document, path: Path, width: float, alt: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt)
    doc_pr.set("title", alt)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_keep_with_next(style) -> None:
    style.paragraph_format.keep_with_next = True


def remove_paragraph_borders(target) -> None:
    if hasattr(target, "_p"):
        p_pr = target._p.get_or_add_pPr()
    else:
        p_pr = target.element.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is not None:
        p_pr.remove(borders)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Source and interpretation note  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(text)


def main() -> None:
    project = load_csv("project_summary.csv")[0]
    sources = load_csv("source_dataset_inventory.csv")
    stages = load_csv("preprocessing_stages.csv")
    encoding = load_csv("encoding_strategy.csv")
    lineage = load_csv("engineered_feature_lineage.csv")
    removed = load_csv("correlation_removed_features.csv")
    insights = load_csv("eda_insights.csv")
    metrics = load_csv("model_metrics.csv")
    ensemble = load_csv("ensemble_results.csv")
    powerbi = load_csv("powerbi_manifest.csv")
    provenance = json.loads((DATA / "provenance.json").read_text(encoding="utf-8"))

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string("1F2933")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title.font.size = Pt(29)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    remove_paragraph_borders(title)

    for style_name, size, before, after in (
        ("Heading 1", 18, 16, 7),
        ("Heading 2", 13.5, 12, 5),
        ("Heading 3", 11.5, 9, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        set_keep_with_next(style)

    caption = styles["Caption"]
    caption.font.name = "Aptos"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(MUTED)
    caption.paragraph_format.space_after = Pt(9)

    if "Document Subtitle" not in styles:
        subtitle = styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        subtitle = styles["Document Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    for sec in doc.sections:
        add_page_number(sec.footer.paragraphs[0])

    # Cover page
    doc.add_paragraph("CREDITSCOPE", style="Subtitle")
    cover_title = doc.add_paragraph("CreditScope Complete Project Documentation", style="Title")
    remove_paragraph_borders(cover_title)
    doc.add_paragraph("From Data Preparation to the Final Web Application", style="Document Subtitle")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(22)
    p.add_run(
        "This document records the complete implemented workflow behind CreditScope. It follows the organized notebook, "
        "the exported research evidence, and the current application code from raw source tables through feature "
        "engineering, modeling, deployment, dashboard analysis, prediction, reporting, and model-assisted explanation."
    )
    cover_rows = [
        ["Primary dataset", "Home Credit Default Risk"],
        ["Applicants", f"{int(project['Applicants']):,}"],
        ["Observed TARGET equals 1", f"{int(project['Target1Count']):,} applicants ({float(project['DefaultRate']):.2%})"],
        ["Engineered features", project["EngineeredFeatures"]],
        ["Final numeric model features", project["NumericFeaturesAfterCorrelation"]],
        ["Application", "Local FastAPI service with a responsive browser dashboard"],
        ["Prepared", "5 October 2026"],
    ]
    add_table(doc, ["Project item", "Recorded value"], cover_rows, [2.2, 4.55], 9.2)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.add_run("Research use statement  ").bold = True
    p.add_run(
        "CreditScope presents historical associations and model scores. It does not make lending decisions, and its "
        "scores are not calibrated real world default probabilities unless a deployed model explicitly demonstrates calibration."
    )
    doc.add_page_break()

    doc.add_heading("Contents", level=1)
    contents = [
        "Project purpose and scope", "Project summary", "End to end workflow", "Source data and lineage",
        "Exploratory data analysis", "Feature engineering", "Data cleaning and preprocessing", "Model development",
        "Model evaluation", "From notebook to deployable assets", "Application architecture", "Website functions",
        "Prediction and explanation logic", "Batch scoring", "Model deployment and registry", "AI assistant",
        "Reporting and exports", "API reference", "Implementation function reference", "Security and deployment boundaries",
        "Validation and testing", "Operation and maintenance", "Known limitations", "Data artifact reference", "Glossary",
    ]
    for item in contents:
        add_bullet(doc, item)
    doc.add_paragraph(
        "The document is organized as a technical project handbook. Early sections explain the analytical workflow; "
        "later sections document the deployed product and its implementation."
    )

    doc.add_heading("Project Purpose and Scope", level=1)
    doc.add_paragraph(
        "CreditScope converts a credit risk data analytics project into a local research application. The notebook establishes "
        "the data preparation, exploratory analysis, feature engineering, validation design, model comparisons, and model exports. "
        "The web application then exposes those assets through portfolio analysis, individual and batch scoring, model deployment, "
        "notebook chart exploration, PDF and data exports, worldwide context, and a constrained AI assistant."
    )
    doc.add_paragraph(
        "The main analytical conclusion is that class imbalance makes accuracy insufficient, historical credit behavior adds useful "
        "signal, and the strongest recorded individual models are gradient boosting models. CatBoost achieved the highest individual "
        "validation ROC AUC at 0.765365. A three model probability blend reached 0.766916, while the final notebook submission used "
        "LightGBM because the test path was complete and consistent for that model."
    )
    add_source_note(doc, "Project facts come from the organized notebook exports, data research tables, README, backend modules, frontend code, and automated tests in this repository.")

    doc.add_heading("Project Summary", level=1)
    summary_rows = [
        ["Dataset", project["Dataset"]],
        ["Training applicants", f"{int(project['Applicants']):,}"],
        ["Class 0", f"{int(project['Target0Count']):,} ({1-float(project['DefaultRate']):.2%})"],
        ["Class 1", f"{int(project['Target1Count']):,} ({float(project['DefaultRate']):.2%})"],
        ["Added engineered features", project["EngineeredFeatures"]],
        ["High missing numeric columns removed", project["HighMissingNumericDropped"]],
        ["Numeric predictors before correlation reduction", project["NumericFeaturesBeforeCorrelation"]],
        ["Final numeric predictors", project["NumericFeaturesAfterCorrelation"]],
        ["CatBoost predictors", f"{project['CatBoostFeatures']} with {project['CatBoostCategoricalFeatures']} native categorical features"],
        ["Models in the complete leaderboard", str(len(metrics))],
        ["Best individual model", f"{project['BestIndividualModel']} with ROC AUC {float(project['BestIndividualROCAUC']):.6f}"],
        ["Best recorded blend", f"ROC AUC {float(project['BestBlendROCAUC']):.6f}"],
    ]
    add_table(doc, ["Measure", "Result"], summary_rows, [3.3, 3.45], 9.0)

    doc.add_heading("End To End Workflow", level=1)
    workflow = [
        "Load and inventory current application, bureau, bureau balance, previous application, installment, test, and submission data.",
        "Audit the target distribution, missingness, distributions, duplicates, infinities, and outliers.",
        "Create five application features and aggregate one to many historical tables to applicant level.",
        "Merge 40 engineered application and history features into the application table.",
        "Remove 38 numerical variables with more than 50 percent missingness, fill remaining numeric missing values with training medians, and preserve categorical missingness as Missing.",
        "Apply shared binary and ordinal mappings, then split into a fully numeric branch and a native categorical CatBoost branch.",
        "Screen categorical variables, frequency encode selected categories, one hot encode retained categories, and remove identifier leakage.",
        "Reduce the numeric matrix from 122 to 107 predictors by removing 15 redundant or highly correlated variables.",
        "Create a stratified 80 percent training and 20 percent validation split with 246008 and 61503 rows.",
        "Train and compare linear, tree boosting, random forest, neural network, and SMOTE experiments using class aware metrics.",
        "Export deployable model packages with the exact feature order, classification threshold, and any required preprocessor.",
        "Import notebook evidence and the Power BI applicant snapshot into the local FastAPI application without executing notebook code.",
        "Expose the results through the browser dashboard, REST API, exports, prediction screens, model registry, and research assistant."
    ]
    for step in workflow:
        add_number(doc, step)

    doc.add_heading("Source Data and Lineage", level=1)
    doc.add_paragraph(
        "The current application table contains one row per application. Historical tables contain repeated rows per applicant, so "
        "they are aggregated before merging. Bureau balance requires two levels of aggregation: monthly rows become bureau credit "
        "summaries, then bureau credits become applicant summaries."
    )
    source_rows = [[r["Dataset"], f"{int(r['Rows']):,}", r["Columns"], r["ProjectUsage"], r["Purpose"]] for r in sources]
    add_table(doc, ["Dataset", "Rows", "Columns", "Use", "Purpose"], source_rows,
              [1.6, 0.95, 0.65, 1.15, 2.4], 7.7, ["left", "right", "center", "center", "left"])
    doc.add_paragraph(
        "The dashboard uses a saved applicant snapshot imported from the Power BI fact_applicants.csv export. It is not a live Power BI connection. "
        "Rerunning the importer refreshes the local snapshot. Original amount units are preserved; the project does not infer a BDT or USD conversion."
    )
    doc.add_paragraph(
        "The source project summary export still contains the earlier project label CrediFuse. The maintained application, code, and user interface use the name CreditScope."
    )

    doc.add_heading("Exploratory Data Analysis", level=1)
    doc.add_heading("Target Distribution", level=2)
    doc.add_paragraph(
        "TARGET records repayment difficulty in the source dataset. Class 1 represents 24825 of 307511 applications, or 8.07 percent. "
        "Because 91.93 percent of records belong to class 0, the project emphasizes ROC AUC, precision recall AUC, precision, recall, and F1 rather than accuracy alone."
    )
    add_picture(doc, DATA / "eda" / "cell-10-0.png", 6.25, "Bar chart showing 282686 class zero applicants and 24825 class one applicants")
    add_caption(doc, "Figure 1  TARGET class distribution from the organized notebook")

    doc.add_heading("Recorded Findings", level=2)
    insight_rows = [[r["Topic"], r["Insight"], r["InterpretationNote"]] for r in insights]
    add_table(doc, ["Topic", "Recorded finding", "Interpretation"], insight_rows, [1.35, 4.25, 1.15], 8.0)
    add_picture(doc, DATA / "eda" / "cell-27-4.png", 6.25, "Bar chart showing observed repayment difficulty falling across increasing age groups")
    add_caption(doc, "Figure 2  Observed repayment difficulty by age group")
    doc.add_paragraph(
        "The age pattern is descriptive, not causal. The notebook also recorded higher observed difficulty for male applicants than female applicants, "
        "lower observed difficulty among higher education applicants, and non linear patterns for income, credit amount, and credit to income ratio."
    )

    doc.add_heading("Feature Engineering", level=1)
    doc.add_paragraph(
        "Forty features were added across six source combinations. Each historical table was aggregated to applicant level before it was merged into the main application table."
    )
    add_picture(doc, DATA / "eda" / "cell-50-0.png", 6.4, "Horizontal bar chart showing engineered feature counts by source dataset")
    add_caption(doc, "Figure 3  Engineered feature counts by source")
    feature_rows = [[str(i), r["Feature"], r["SourceDataset"], r["Definition"]] for i, r in enumerate(lineage, 1)]
    add_table(doc, ["No", "Feature", "Source", "Definition"], feature_rows, [0.42, 2.15, 1.55, 2.63], 7.6,
              ["center", "left", "left", "left"])
    doc.add_paragraph(
        "Application features include income and credit quintiles, credit to income, age in years, and age bands. Bureau features summarize "
        "counts, active and overdue ratios, overdue severity, debt totals, and debt burden. Bureau balance features summarize monthly history. "
        "Previous application and installment features capture earlier decisions, average amounts, late payment, payment delay, and underpayment."
    )

    doc.add_heading("Data Cleaning and Preprocessing", level=1)
    doc.add_heading("Recorded Pipeline Stages", level=2)
    stage_rows = [[r["StageOrder"], r["Stage"], f"{int(float(r['Rows'])):,}", r["Columns"], r["Action"]] for r in stages]
    add_table(doc, ["Step", "Stage", "Rows", "Columns", "Action"], stage_rows,
              [0.48, 1.85, 0.85, 0.7, 2.87], 7.7, ["center", "left", "right", "center", "left"])

    doc.add_heading("Missing Values and Outliers", level=2)
    add_bullet(doc, "Numerical columns with more than 50 percent missingness were removed after the full engineered table was assembled. The source run removed 38 columns, including EXT_SOURCE_1 and highly missing building variables.")
    add_bullet(doc, "Remaining numerical missing values were filled with medians learned from the training data. The importer also stores a training default for every deployed feature.")
    add_bullet(doc, "Low cardinality numerical features use the training mode as the application default so binary and ordinal values remain valid.")
    add_bullet(doc, "Categorical missing values were represented as Missing before categorical processing.")
    add_bullet(doc, "The IQR procedure detected outliers but did not globally remove or cap them. The explicit domain correction clipped negative aggregated bureau debt to zero.")
    add_bullet(doc, "Training medians, modes, and quantile boundaries are reused at inference. Test or manual-case values do not refit preprocessing.")

    doc.add_heading("Categorical Processing", level=2)
    encoding_rows = [[r["Feature"], r["Strategy"], r["Detail"]] for r in encoding]
    add_table(doc, ["Feature", "Method", "Implementation"], encoding_rows, [2.15, 1.75, 2.85], 7.7)
    doc.add_paragraph(
        "Chi square tests and Cramer's V supported the categorical screening step. NAME_INCOME_TYPE, WALLSMATERIAL_MODE, and EMERGENCYSTATE_MODE were retained for one hot encoding. "
        "Six weaker categorical variables were removed from the numeric branch. Rare income categories were grouped into Other."
    )

    doc.add_heading("Redundancy Reduction", level=2)
    doc.add_paragraph(
        "The numeric branch contained 122 predictors after encoding and identifier removal. Fifteen variables were removed because of duplicate information or high correlation, producing the final 107 feature matrix. "
        "The clearest duplicate was DAYS_BIRTH versus the derived AGE_YEARS, with an absolute correlation of 1.0."
    )
    removed_rows = [[r["Feature"], r["Decision"]] for r in removed]
    add_table(doc, ["Removed feature", "Decision"], removed_rows, [2.4, 4.35], 8.4)

    doc.add_heading("Model Development", level=1)
    doc.add_heading("Validation Design", level=2)
    doc.add_paragraph(
        "The final numeric matrix was split into 246008 training rows and 61503 validation rows using a stratified 80 to 20 split. "
        "Scaling was fitted on training data for scale sensitive models. The original validation set was retained for evaluation after any training-only SMOTE resampling."
    )
    doc.add_heading("Modeling Methods", level=2)
    model_method_rows = [
        ["Logistic Regression top 40", "StandardScaler; max_iter 1000; balanced class weights; Random Forest selected top 40 features"],
        ["Random Forest top 40", "200 trees; balanced class weights; used both as a baseline and as the feature-importance selector"],
        ["Logistic Regression 107", "StandardScaler; max_iter 1000; balanced class weights; all final numeric predictors"],
        ["LightGBM", "1500 estimators; learning rate 0.03; 31 leaves; row and column subsampling 0.8; L1 and L2 regularization 0.1; balanced classes; early stopping after 100 rounds"],
        ["CatBoost", "3000 iterations; learning rate 0.03; depth 8; Logloss; AUC evaluation; automatic balanced class weights; 11 native categorical variables; early stopping 150 rounds"],
        ["XGBoost", "2000 estimators; learning rate 0.03; depth 8; min child weight 5; subsampling 0.85; gamma and L1 0.1; L2 1.0; positive-class weighting; early stopping 100 rounds"],
        ["MLP", "Median imputer and StandardScaler pipeline; hidden layers 128 64 32; ReLU and Adam; batch size 1024; early stopping; balanced sample weights when supported"],
        ["Deep and Cross Network", "Median imputer and StandardScaler; three cross layers; deep branch 256 128 64 with batch normalization and dropout; weighted BCE loss; AdamW; learning-rate reduction and early stopping"],
        ["Logistic Regression with SMOTE", "Top 40 scaled features; SMOTE applied only to training data; validation remained in the original distribution"],
        ["MLP with SMOTE", "Top 40 scaled features; hidden layers 64 32; Adam; early stopping; training data resampled only"],
        ["DCN with SMOTE", "Notebook name for a deeper scikit-learn MLP with layers 256 128 64 32; it is distinct from the PyTorch Deep and Cross Network"],
    ]
    add_table(doc, ["Method", "Recorded configuration"], model_method_rows, [2.0, 4.75], 8.0)
    doc.add_paragraph(
        "The LightGBM sensitivity checks increased num_leaves to 63 and removed class weighting. Their ROC AUC values, 0.763842 and 0.763984, remained below the balanced 31 leaf baseline at 0.764678."
    )

    doc.add_heading("Model Evaluation", level=1)
    doc.add_paragraph(
        "ROC AUC measures ranking quality across thresholds. PR AUC is especially useful for the 8.07 percent positive class. Precision measures how many flagged cases are positive, recall measures how many positives are found, and F1 balances precision and recall. "
        "Accuracy can look high when a model mostly predicts the majority class, as shown by the Random Forest baseline."
    )
    ranking = sorted(metrics, key=lambda r: float(r["ROC_AUC"]), reverse=True)
    ranking_rows = [[r["Model"], r["Group"], f"{float(r['ROC_AUC']):.4f}", f"{float(r['PR_AUC']):.4f}"] for r in ranking]
    add_table(doc, ["Model", "Group", "ROC AUC", "PR AUC"], ranking_rows, [2.85, 1.45, 1.15, 1.15], 8.2,
              ["left", "left", "right", "right"])
    classification_rows = [[r["Model"], f"{float(r['Accuracy']):.4f}", f"{float(r['Precision']):.4f}", f"{float(r['Recall']):.4f}", f"{float(r['F1']):.4f}"] for r in ranking]
    add_table(doc, ["Model", "Accuracy", "Precision", "Recall", "F1"], classification_rows, [2.85, 0.98, 0.98, 0.98, 0.98], 7.9,
              ["left", "right", "right", "right", "right"])
    add_picture(doc, DATA / "eda" / "cell-127-2.png", 6.45, "Horizontal bar chart comparing validation ROC AUC for eleven models")
    add_caption(doc, "Figure 4  Validation ROC AUC comparison from the notebook")
    doc.add_paragraph(
        "CatBoost was the strongest recorded individual model at 0.765365 ROC AUC and 0.257852 PR AUC. LightGBM was close at 0.764678 and 0.254421. "
        "The best recorded blend used 0.4 LightGBM, 0.5 CatBoost, and 0.1 XGBoost and reached 0.766916 ROC AUC and about 0.25927 PR AUC."
    )
    doc.add_paragraph(
        "The top 40 Random Forest illustrates why accuracy was not used alone. It reached 0.9191 accuracy but only 0.0024 recall, with 12 true positives and 4953 false negatives at its recorded decision rule."
    )

    doc.add_heading("From Notebook to Deployable Assets", level=1)
    doc.add_paragraph(
        "The importer reads saved notebook outputs and data exports without executing notebook code. It stores the applicant snapshot, copies small research CSV files, calculates training-only defaults and quantile boundaries, extracts 45 saved notebook charts, exports the final model leaderboard, and records provenance."
    )
    for item in [
        "data/applicants.pkl stores the full applicant-level dashboard snapshot.",
        "data/schema.json stores the exact 107-feature order, per-feature training defaults and methods, bin boundaries, and training row count.",
        "data/gallery.json and data/eda store chart metadata and the 45 preserved notebook images.",
        "data/notebook_context.json stores notebook markdown used by the research assistant.",
        "data/research stores the exported analysis, modeling, Power BI, and provenance tables.",
        "Optional model import copies the original MLP pipeline and DCN checkpoint into the local model registry.",
    ]:
        add_bullet(doc, item)
    doc.add_paragraph(
        "The final notebook submission used LightGBM for 48744 test rows. It verified the feature order, output shape, unique applicant identifiers, absence of missing values, and scores between zero and one. The source notebook recorded a mean test score of 0.319657."
    )

    doc.add_heading("Application Architecture", level=1)
    architecture_rows = [
        ["Browser interface", "frontend/index.html, app.js, and styles.css", "Hash routed single page interface, charts, forms, exports, accessible dialogs, and state management"],
        ["API service", "backend/main.py", "FastAPI routes, request validation, upload handling, security boundaries, chat integration, and static files"],
        ["Portfolio analytics", "backend/analytics.py", "Cached applicant snapshot, six filters, grouped rates, KPIs, age and income matrix, and research table loading"],
        ["Feature construction", "backend/features.py", "Twelve input schema, training defaults, derived values, native categorical support, and exact model feature order"],
        ["Model registry", "backend/models.py", "Trusted artifact loading, model adapters, smoke tests, cached inference, thresholds, and local sensitivities"],
        ["Batch scoring", "backend/batch.py", "CSV parsing, header matching, row conversion, feature pass through, limits, and one hour detail sessions"],
        ["Neural architecture", "backend/dcn.py", "PyTorch CrossLayer and DeepCrossNetwork matching the saved notebook checkpoint"],
        ["Report generation", "backend/reporting.py", "Two page filtered executive dashboard PDF using ReportLab"],
        ["Research data", "data directory", "Applicant snapshot, schema, global indicators, notebook text, charts, and exported research tables"],
        ["Local model storage", "storage/models and storage/archived", "Active validated models and recoverable archived removals"],
    ]
    add_table(doc, ["Layer", "Implementation", "Responsibility"], architecture_rows, [1.35, 2.25, 3.15], 8.0)
    doc.add_paragraph(
        "The browser sends requests to the local FastAPI server. The server builds exact ordered model inputs from the supplied applicant or CSV row, applies the saved model adapter, and returns the score, threshold decision, assumptions, and local sensitivity checks."
    )

    doc.add_heading("Website Functions", level=1)
    doc.add_heading("Executive Overview", level=2)
    for item in [
        "Analyzes all 307511 applicants or a filtered cohort using age group, gender, education, income type, bureau overdue, and debt to income filters.",
        "Shows applicant count, observed difficulty rate, average income, average credit, average age, total credit exposure, recorded difficulty exposure, and average credit to income ratio.",
        "Displays outcome composition, age group rates, an age and income risk matrix, debt burden, income type, and education views.",
        "Allows chart interaction to apply age filters and supports a one-click reset.",
        "Exports the current portfolio summary as JSON and generates a two page PDF using the active filters.",
        "Provides a simplified expected loss scenario using observed risk rate times a user-selected LGD times credit exposure."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Worldwide Statistics", level=2)
    for item in [
        "Shows a Robinson projection map and a country selector for the saved global indicator snapshot.",
        "Presents bank non-performing loans, GDP growth, inflation, unemployment, lending rate, bank capital to assets, and private-sector credit to GDP with the recorded reporting year.",
        "Shows annual NPL movement where a previous value is available.",
        "Provides BIS worldwide banking exposure and a transparent illustrative loss scenario using the median NPL ratio of available countries and a 40 percent LGD assumption.",
        "Links metric definitions and official sources while keeping national indicators separate from applicant predictions."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Individual Prediction", level=2)
    for item in [
        "Accepts 12 inputs covering income, credit, annuity, contract, age, employment, education, income type, car ownership, external scores 2 and 3, and late payment rate.",
        "Lets the user choose any deployed model and validates all inputs before scoring.",
        "Displays the class 1 model score, saved decision threshold, flagged status, number of training defaults, and all model feature values with their origin.",
        "Calculates one-input-at-a-time local sensitivity by replacing a supplied input with its training reference and recomputing dependent features.",
        "Compares the same case across all deployed models.",
        "Exports the full case result as JSON and adds an applicant expected-loss scenario.",
        "Activates applicant-specific chat context after a successful score."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Batch CSV Check", level=2)
    for item in [
        "Accepts UTF 8 CSV, tab, semicolon, or pipe delimited text saved with a .csv filename, with unique headers and up to 50000 rows or 60 MB.",
        "Recognizes human-readable form names, raw Home Credit names, and direct deployed-model feature names.",
        "Scores valid rows in one model batch, reports row errors separately, and ranks results by model score.",
        "Shows exposure-weighted score, flagged count, high-risk exposure, expected-loss scenario, matched column counts, distinct score count, and training defaults used.",
        "Paginates results in groups of 100, exports a formula-safe CSV, and warns when every row receives the same score.",
        "Allows any valid row to be reopened for its full feature evidence, sensitivities, JSON report, model comparison, and applicant chat.",
        "Keeps at most two in-memory batch sessions. A detail session expires after one hour or an application restart."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Explore The Data", level=2)
    for item in [
        "Provides all 45 preserved notebook charts with plot title, source cell, section, category, and enlarged inspection.",
        "Filters the gallery by exploratory analysis, model evaluation, pipeline, plot group, and search text.",
        "Shows the 14 preprocessing stages and the 11-model leaderboard.",
        "Summarizes best individual ROC AUC and recall while keeping model metrics tied to the validation set."
    ]:
        add_bullet(doc, item)

    doc.add_heading("AI Assistant", level=2)
    for item in [
        "Supports a local Ollama assistant and a Gemini cloud assistant with separate browser-memory conversation histories.",
        "General chat receives aggregate dashboard context, notebook notes, exported findings, reported metrics, and the live model registry.",
        "Applicant chat receives a server-verified score, threshold, input origins, top local sensitivities, and a historical peer cohort defined by age within five years and income within 25 percent.",
        "Renders a limited Markdown subset safely after escaping HTML.",
        "Limits conversations to 12 messages, user content to 5000 characters, output to 1200 tokens, and concurrent chat calls to two.",
        "Does not train models, change predictions, make lending decisions, or treat local sensitivities as causal explanations."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Deploy Model", level=2)
    for item in [
        "Lists locally deployed models, feature counts, thresholds, and reported validation ROC AUC when supplied.",
        "Downloads a starter metadata template.",
        "Accepts trusted scikit-learn, PyTorch DCN checkpoint, TorchScript, LightGBM, CatBoost, and optional Keras artifacts.",
        "Validates extension, metadata, threshold, exact ordered feature names, categorical features, preprocessing requirements, model output shape, finite values, and the zero-to-one output range.",
        "Performs a reference-applicant smoke inference before the upload becomes selectable.",
        "Removes a model from active prediction while moving its directory to a recoverable local archive."
    ]:
        add_bullet(doc, item)

    doc.add_heading("Prediction and Explanation Logic", level=1)
    input_rows = [
        ["age", "19 to 100", "AGE_YEARS and AGE_GROUP"], ["income", "1000 to 100000000", "AMT_INCOME_TOTAL, income group, credit ratios"],
        ["credit", "1000 to 100000000", "AMT_CREDIT, credit group, credit to income"], ["annuity", "1 to 10000000", "AMT_ANNUITY"],
        ["employment_years", "0 to 85 and cannot start before age 14", "DAYS_EMPLOYED as negative days"],
        ["education", "Five ordered categories", "NAME_EDUCATION_TYPE level"], ["income_type", "Five categories", "One hot income fields or native category"],
        ["contract_type", "Cash or revolving", "NAME_CONTRACT_TYPE binary"], ["own_car", "Yes or no", "FLAG_OWN_CAR"],
        ["external_score_2", "Optional zero to one", "EXT_SOURCE_2 or training default"], ["external_score_3", "Optional zero to one", "EXT_SOURCE_3 or training default"],
        ["late_payment_rate", "Optional zero to one", "LATE_PAYMENT_RATE or training default"],
    ]
    add_table(doc, ["Input", "Validation", "Model effect"], input_rows, [1.55, 2.5, 2.7], 8.0)
    doc.add_paragraph(
        "Credit to income, age group, income group, credit group, bureau debt to income, and debt bands are recalculated for each case. "
        "Historical details that cannot be known from the 12 fields retain their training medians or modes. The returned assumptions table exposes every entered, derived, CSV-supplied, and training-default feature."
    )
    doc.add_paragraph(
        "For a manual case, each sensitivity replaces one of the 12 user inputs with its reference value and rebuilds dependent features. For a batch detail, each supplied model feature is replaced separately with the model reference. "
        "The score difference is local and non-additive. It is not SHAP, causal attribution, or proof that changing a field would change an actual credit outcome."
    )

    doc.add_heading("Batch Scoring", level=1)
    doc.add_paragraph(
        "The parser removes a UTF 8 byte-order mark, detects comma, tab, semicolon, or pipe delimiters, trims headers, rejects duplicate or missing headers, and rejects rows with more values than the header. "
        "It converts raw DAYS_BIRTH to age, negative DAYS_EMPLOYED to employment years, encoded education and contract values to labels, and common binary strings to booleans."
    )
    doc.add_paragraph(
        "Direct model feature columns take precedence when present. Numeric values must be finite. CatBoost categorical values pass through as strings. Selected raw categories such as gender, ownership, contract, education, occupation, and organization are converted using the same rules or frequency maps used by the application."
    )

    doc.add_heading("Model Deployment and Registry", level=1)
    package_rows = [
        ["Scikit-learn", ".joblib or .pkl", "Model with predict_proba; binary classes 0 and 1; pipeline or separate preprocessor"],
        ["PyTorch DCN checkpoint", ".pth or .pt", "Matching checkpoint architecture, exact feature order, threshold, and saved preprocessor"],
        ["PyTorch TorchScript", ".pt", "Scripted model, exact feature order, threshold, and preprocessor"],
        ["LightGBM", ".txt", "Native Booster file and ordered feature list"],
        ["CatBoost", ".cbm", "Native model, ordered features, and declared categorical feature names"],
        ["Keras", ".h5 or .keras", "Complete safe-load model, matching input shape, optional external preprocessor, and optional TensorFlow runtime"],
    ]
    add_table(doc, ["Runtime", "Artifact", "Requirements"], package_rows, [1.6, 1.15, 4.0], 8.1)
    doc.add_paragraph(
        "Each upload must include a classification threshold strictly between zero and one and an ordered feature list. The server calculates a SHA 256 digest, writes clean metadata, and caches the validated model adapter. "
        "Serialized pickle and joblib files can execute code, so the interface requires an explicit trusted-file confirmation."
    )
    doc.add_paragraph(
        "SMOTE is applied during training only. New applicants are never resampled. The export helper saves a fitted estimator, fitted preprocessor, exact feature order, threshold, calibration flag, and input-space metadata without retraining the model."
    )

    doc.add_heading("AI Assistant", level=1)
    doc.add_paragraph(
        "Ollama is restricted to a localhost or loopback URL. Gemini requires GEMINI_API_KEY and sends the conversation and any selected applicant context to the configured Google model. "
        "The assistant system instructions require source naming, separation of validation metrics from live model capability, clear score caveats, and refusal to invent missing results or make lending decisions."
    )
    doc.add_paragraph(
        "A selected applicant brief is recomputed on the server rather than trusted from browser text. It includes the model score, threshold, decision label, assumptions count, top 12 local sensitivities, portfolio rate, and a comparable historical cohort."
    )

    doc.add_heading("Reporting and Exports", level=1)
    export_rows = [
        ["Portfolio JSON", "Current filters and all dashboard summary data"],
        ["Dashboard PDF", "Two-page filtered executive report with KPIs, outcome table, age and debt charts, income and education tables, and age-income matrix"],
        ["Case JSON", "Prediction score, threshold, decision, inputs, assumptions, feature origins, local sensitivities, caveats, case identifier, and timestamp"],
        ["Batch CSV", "Row, applicant identifier, exposure, model score, flag, current expected-loss scenario, and defaults used; spreadsheet formulas are neutralized"],
        ["Metadata template", "Starter model metadata including 107 feature names and required deployment fields"],
        ["Interactive API reference", "Local Swagger UI at the /docs route using bundled assets"],
    ]
    add_table(doc, ["Output", "Contents"], export_rows, [1.6, 5.15], 8.3)

    doc.add_heading("API Reference", level=1)
    endpoints = [
        ["GET", "/api/health", "Runtime readiness, assistant configuration, and admin-token status"],
        ["GET", "/api/config", "Dashboard filters, reference applicant, provenance, global indicators, and health"],
        ["GET", "/api/dashboard", "Portfolio metrics and grouped rates with optional six-field filtering"],
        ["GET", "/api/dashboard/report", "Filtered dashboard PDF; rejects empty cohorts"],
        ["GET", "/api/eda", "Plot catalog, EDA insights, preprocessing stages, and model metrics"],
        ["GET", "/api/models", "Active validated model registry"],
        ["GET", "/api/models/template", "Starter deployment metadata"],
        ["POST", "/api/predict", "Score one validated Applicant with one deployed model"],
        ["POST", "/api/compare", "Score one Applicant across all active models"],
        ["POST", "/api/batch/predict", "Parse and score an uploaded applicant CSV"],
        ["POST", "/api/batch/detail", "Recreate full evidence for one retained batch row"],
        ["POST", "/api/models", "Validate and deploy a trusted model package"],
        ["DELETE", "/api/models/{model_id}", "Remove an active model and move it to the local archive"],
        ["POST", "/api/chat", "Ask Ollama or Gemini with project and optional applicant context"],
        ["GET", "/", "Serve the main CreditScope application"],
        ["GET", "/docs", "Serve the local interactive API reference"],
    ]
    add_table(doc, ["Method", "Path", "Purpose"], endpoints, [0.72, 2.35, 3.68], 8.2, ["center", "left", "left"])

    doc.add_heading("Implementation Function Reference", level=1)
    doc.add_paragraph(
        "This reference covers the named application functions that implement the workflow. Small formatting helpers are included because they directly support the final interface and exports."
    )
    backend_functions = [
        ["analytics.py", "applicants", "Load and cache the applicant snapshot"], ["analytics.py", "records", "Convert a pandas frame to JSON-safe records"],
        ["analytics.py", "research", "Load an exported research CSV"], ["analytics.py", "group", "Calculate applicant count, defaults, and rate for a field"],
        ["analytics.py", "dashboard", "Apply filters and assemble every portfolio metric and chart dataset"], ["analytics.py", "filter_options", "Build the six filter menus"],
        ["features.py", "Applicant employment_check", "Reject employment durations that would begin before age 14"], ["features.py", "schema", "Load feature order, defaults, methods, and bins"],
        ["features.py", "bucket", "Apply right-closed notebook bin semantics with range clamping"], ["features.py", "make_features", "Build ordered numeric model inputs and assumptions"],
        ["features.py", "category_mode", "Find and cache a training category mode"], ["features.py", "category_frequency", "Find and cache category frequencies"],
        ["features.py", "make_model_features", "Build exact numeric or native-categorical model rows with CSV overrides"], ["features.py", "reference_case", "Create the training-reference applicant"],
        ["models.py", "read_metadata", "Read one deployed model's metadata"], ["models.py", "registry", "List usable models while skipping incomplete directories"],
        ["models.py", "model_folder", "Validate a model identifier and locate its directory"], ["models.py", "load", "Load, validate, adapt, smoke test, and cache a supported model"],
        ["models.py", "transform", "Apply a saved imputer-scaler dictionary or generic transformer"], ["models.py", "predict_case", "Score a case and create thresholds, assumptions, sensitivities, labels, and caveats"],
        ["batch.py", "remember", "Store up to two one-hour batch sessions"], ["batch.py", "selected", "Retrieve a retained row with expiration and range checks"],
        ["batch.py", "parse_csv", "Decode, detect delimiter, validate headers, and enforce file limits"], ["batch.py", "recognized_columns", "Identify uploaded fields that can affect the selected model"],
        ["batch.py", "applicant_from_row", "Convert form, raw, or encoded CSV values to an Applicant"], ["batch.py", "model_values", "Pass direct feature matches through with required encodings"],
        ["dcn.py", "CrossLayer forward", "Apply one explicit feature-cross layer"], ["dcn.py", "DeepCrossNetwork forward", "Combine cross and deep branches for one logit"],
        ["reporting.py", "format helpers", "Format counts, compact values, percentages, and safe paragraphs"], ["reporting.py", "chart helpers", "Create age, debt, and age-income visuals for the PDF"],
        ["reporting.py", "build_dashboard_report", "Generate the complete filtered two-page PDF in memory"], ["main.py", "boundaries", "Enforce trusted hosts, same-origin writes, and security response headers"],
        ["main.py", "admin", "Require the configured admin token for model and AI operations"], ["main.py", "health and config", "Expose startup, assistant, filter, provenance, and global context"],
        ["main.py", "dashboard and dashboard_report", "Serve live portfolio analysis and PDF output"], ["main.py", "eda and model_list", "Serve notebook evidence and the active model list"],
        ["main.py", "batch_predict and batch_detail", "Score files and reconstruct detailed rows"], ["main.py", "predict and compare", "Score one case with one or all models"],
        ["main.py", "save_upload and deploy", "Stream uploads, validate packages, run smoke inference, and register models"], ["main.py", "delete", "Archive an active model"],
        ["main.py", "chat_instructions", "Create constrained project-grounded assistant instructions"], ["main.py", "applicant_brief", "Create verified selected-case chat evidence"],
        ["main.py", "ollama_url and chat", "Restrict local Ollama URLs and call Ollama or Gemini"], ["import_project.py", "write and import workflow", "Build the local data, schema, gallery, notebook context, provenance, and optional model snapshot"],
        ["export_smote_models.py", "export_model", "Package a fitted SMOTE model, preprocessor, feature order, threshold, and metadata"],
    ]
    add_table(doc, ["Module", "Function", "Purpose"], backend_functions, [1.35, 2.05, 3.35], 7.7)

    frontend_functions = [
        ["Formatting and safety", "$ $$ esc number compact pct money", "DOM selection, HTML escaping, and display formatting"],
        ["Network and files", "api download toast", "Call the API, report errors, download generated content, and show status"],
        ["Shared interface", "heading stat panel bars matrix", "Create page headings, KPIs, panels, bar charts, and the segment matrix"],
        ["Metric guidance", "infoButton bindInfoButtons", "Open accessible explanations and official-source links"],
        ["Worldwide context", "robinsonPoint globalRiskMarkup defaultRiskPulseMarkup bindGlobalRisk renderGlobal", "Place countries on the map and render macro and banking context"],
        ["Loss scenario", "lossPanelMarkup bindLossPanel portfolioLossMarkup", "Render and update PD times LGD times EAD scenarios"],
        ["Portfolio", "renderOverview refreshDashboard", "Build filters, fetch cohorts, render charts, and export current results"],
        ["Manual case", "field readCase renderPredict", "Build the 12-field form, parse values, run prediction, and invalidate stale scores"],
        ["Results", "selectChatCase resultMarkup bindResult", "Render score evidence, downloads, cross-model comparison, and case chat"],
        ["Batch", "renderBatch renderBatchResults", "Upload CSVs, show progress, paginate, export, inspect rows, and open case evidence"],
        ["Notebook evidence", "renderEDA renderEDABody researchMetrics renderGallery", "Render stages, leaderboard, filters, search, and enlarged charts"],
        ["Assistant", "currentAssistant formattedAnswer renderCaseChat renderAssistant renderMessages sendChat", "Switch providers, safely render replies, manage histories, and send context"],
        ["Model operations", "adminInput bindAdmin renderModels refreshModels", "Collect local admin token, deploy, remove, and refresh models"],
        ["Navigation", "currentPage navigate start", "Initialize config and models, route pages, handle errors, and update accessibility state"],
    ]
    add_table(doc, ["Area", "Functions", "Purpose"], frontend_functions, [1.35, 2.75, 2.65], 7.6)

    doc.add_heading("Security and Deployment Boundaries", level=1)
    security_points = [
        "The current product is a local research application, not an internet-hardened multi-user service.",
        "Only localhost, loopback, and test hosts are accepted by default. Cross-origin POST, DELETE, PUT, and PATCH requests are rejected.",
        "Responses add nosniff, same-origin referrer, frame-denial, and a restrictive content security policy.",
        "CREDITSCOPE_ADMIN_TOKEN can protect model mutations and AI chat. Remote model management is rejected without the token.",
        "Upload extensions, file sizes, feature order, feature count, categorical declarations, threshold, preprocessing, binary classes, and output ranges are validated.",
        "Trusted serialized model files remain a code-execution risk. Remote hosting should isolate model loading and add authentication, authorization, HTTPS, rate limits, monitoring, backups, and a borrower-data policy.",
        "The application does not delete original notebook model files. Removing a deployed model moves only the local registry copy into storage/archived.",
        "Gemini is a cloud service and receives applicant context when the user selects a case. Ollama keeps assistant processing on the local computer after the model has been downloaded."
    ]
    for item in security_points:
        add_bullet(doc, item)

    doc.add_heading("Validation and Testing", level=1)
    doc.add_paragraph(
        "The current repository test suite was run on 5 October 2026 using the existing project virtual environment and a workspace-local temporary directory. Result: 16 tests passed. Seven warnings were reported, including a FastAPI TestClient deprecation, a future TorchScript load deprecation, and StandardScaler feature-name warnings. No test failed."
    )
    test_rows = [
        ["Dashboard", "Full totals, six filters, global indicators, filtered cohorts, empty cohorts, and unknown filters"],
        ["Reporting", "Two-page PDF generation with active filters and empty-cohort rejection"],
        ["Features", "Exact 107-feature order, derived ratios and bands, optional fields, finite values, and training defaults"],
        ["Inference", "MLP parity, DCN loading, score bounds, thresholds, and removal behavior"],
        ["Security", "Input rejection, cross-origin writes, untrusted hosts, admin-token protection, and bad model uploads"],
        ["Research", "45 notebook plots, 11 leaderboard rows, confusion matrix coverage, and asset availability"],
        ["Batch", "Raw and encoded CSVs, tab delimiters, varying scores, row details, and selected-case context"],
        ["Assistants", "Both providers receive the verified selected applicant context"],
        ["Deployment", "LightGBM, CatBoost, and TorchScript package deployment and scoring"],
    ]
    add_table(doc, ["Area", "Verified behavior"], test_rows, [1.35, 5.4], 8.2)

    doc.add_heading("Operation and Maintenance", level=1)
    doc.add_heading("Start The Application", level=2)
    for index, step in enumerate([
        "Install Git, Git LFS, and Python 3.13.",
        "Clone the repository, run git lfs pull, create a virtual environment, and install requirements.txt.",
        "Copy .env.example to .env and configure Ollama, Gemini, or an admin token only when needed.",
        "Run start.bat and open http://127.0.0.1:8000. Keep the launcher terminal open.",
        "Open http://127.0.0.1:8000/docs for the interactive API reference.",
    ], 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(f"{index}.  ").bold = True
        paragraph.add_run(step)
    doc.add_heading("Refresh Project Data", level=2)
    doc.add_paragraph(
        "Run scripts/import_project.py only when the saved snapshot must be rebuilt from the organized notebook and Home Credit source folder. "
        "CREDITSCOPE_SOURCE can point to a different source directory. The importer does not execute notebook cells and does not modify the source files."
    )
    doc.add_heading("Configure Assistants", level=2)
    add_bullet(doc, "Ollama defaults to http://127.0.0.1:11434 with model gemma3:4b. Start Ollama and download the chosen model before using local chat.")
    add_bullet(doc, "Gemini requires GEMINI_API_KEY. GEMINI_MODEL defaults to gemini-3.5-flash-lite and can be changed to another supported identifier.")

    doc.add_heading("Known Limitations", level=1)
    limitations = [
        "TARGET is the source dataset label for repayment difficulty. It is not a newly defined future default horizon.",
        "Class-weighted and SMOTE-trained model scores are not automatically calibrated population default probabilities.",
        "Manual cases provide only 12 fields. Most historical values are independently filled from training medians or modes, so the resulting history profile may not represent a realistic joint applicant history.",
        "Validation metrics on complete model rows do not measure performance on median-filled manual cases.",
        "One-at-a-time sensitivity is local, non-additive, and non-causal. It should not be described as SHAP.",
        "The expected-loss display is a simplified score or rate times LGD times exposure scenario. It is not an IFRS 9 calculation, and the source currency is not converted even though the interface uses a fixed dollar symbol.",
        "Worldwide indicators are a saved snapshot with different reporting years by country. National conditions do not determine an individual applicant result.",
        "The app has no user accounts and is not ready for unrestricted internet exposure.",
        "The source notebook did not complete a reliable CatBoost test matrix, so the final submission remained LightGBM-only.",
        "The notebook DCN with SMOTE is a deep MLP experiment, not the deployed PyTorch Deep and Cross Network architecture.",
    ]
    for item in limitations:
        add_bullet(doc, item)

    doc.add_heading("Data Artifact Reference", level=1)
    artifact_rows = [[r["PowerBITable"], r["File"], r["Description"], r["RecommendedUse"]] for r in powerbi]
    add_table(doc, ["Table", "File", "Description", "Recommended use"], artifact_rows, [1.35, 1.8, 2.1, 1.5], 7.3)
    doc.add_paragraph(
        "Large curve and pairwise files support interactive or external analysis. The browser application directly consumes the applicant snapshot, schema, gallery, notebook context, provenance, global risk snapshot, EDA insights, preprocessing stages, and model metrics."
    )

    doc.add_heading("Glossary", level=1)
    glossary = [
        ["TARGET", "Historical repayment difficulty label in the Home Credit source data"],
        ["ROC AUC", "Area under the receiver operating characteristic curve; measures ranking across thresholds"],
        ["PR AUC", "Area under the precision recall curve; useful when the positive class is uncommon"],
        ["Precision", "Share of flagged cases that belong to class 1"],
        ["Recall", "Share of class 1 cases that the decision rule flags"],
        ["F1", "Harmonic mean of precision and recall"],
        ["SMOTE", "Synthetic Minority Oversampling Technique applied to training data only"],
        ["DCN", "Deep and Cross Network in the deployed PyTorch architecture; the notebook also used the label for one deeper SMOTE MLP"],
        ["Model score", "Class 1 output produced by a model; not necessarily a calibrated probability"],
        ["Threshold", "Saved score boundary used to assign the flagged or not-flagged label"],
        ["Local sensitivity", "Score change after replacing one supplied value with a reference value"],
        ["PD", "Probability of default or, in this application, the displayed risk score or observed rate used in a scenario"],
        ["LGD", "Loss given default assumption selected by the user"],
        ["EAD", "Exposure at default proxy based on requested credit or portfolio credit"],
    ]
    add_table(doc, ["Term", "Meaning in CreditScope"], glossary, [1.45, 5.3], 8.4)

    core = doc.core_properties
    core.title = "CreditScope Complete Project Documentation"
    core.subject = "End to end data analytics and web application documentation"
    core.author = "CreditScope Project"
    core.keywords = "CreditScope, Home Credit, data preprocessing, feature engineering, machine learning, FastAPI, dashboard"
    core.comments = "Generated from the current repository and exported notebook evidence."

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()

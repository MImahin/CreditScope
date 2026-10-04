from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "DA_Project_Home_Credit_Notebook_Grounded_Report.docx"
DEFAULT_OUTPUT = ROOT / "DA_Project_Home_Credit_Complete_Report_With_Web_Application.docx"


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def replace_paragraph_text(paragraph, text: str):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_column_width(cell, width):
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width.inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_table(doc, headers, rows, widths, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_column_width(cell, Inches(widths[index]))
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, "Aptos", font_size, True)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_column_width(cell, Inches(widths[index]))
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "F7F9FB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, "Aptos", font_size)
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    return table


def add_code_block(doc, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    set_column_width(cell, Inches(7.1))
    set_cell_shading(cell, "F3F4F6")
    set_cell_margins(cell, top=110, start=110, bottom=110, end=110)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, "Consolas", 8.2)
    set_table_borders(table, "B7BEC7")
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.keep_together = True
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.keep_together = True
        p.add_run(item)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.widow_control = True
    return p


def update_front_matter(doc):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Dataset details"):
            replace_paragraph_text(
                paragraph,
                "Dataset details • preprocessing • EDA • feature engineering • model training • web application • deployment",
            )
        elif paragraph.text.startswith("Prepared from notebook state"):
            replace_paragraph_text(
                paragraph,
                "Prepared from the captured notebook state and the implemented CreditScope repository. Notebook cell references use 0-based cell order; application sections describe the current source code and tested local runtime.",
            )
    source_cell = doc.tables[0].cell(0, 0)
    replace_paragraph_text(
        source_cell.paragraphs[0],
        "SOURCE-OF-TRUTH RULE: Notebook pipeline steps, feature counts, metrics, and analytical claims are taken from mainul.ipynb and mainul_2nd.ipynb. The web application extension is taken from the implemented CreditScope source code, README, data assets, and tests. Any notebook inconsistency or deployment boundary is stated explicitly.",
    )
    header = doc.sections[0].header.paragraphs[0]
    replace_paragraph_text(header, "Notebook and application report")

    paragraphs = doc.paragraphs
    position = next(
        index for index, paragraph in enumerate(paragraphs)
        if paragraph.style.name == "Heading 1" and paragraph.text.startswith("1. Project data inventory")
    )
    page_break = paragraphs[position - 1]
    entries = [
        "13. Productization from notebook to application",
        "14. CreditScope architecture and execution flow",
        "15. Website pages and user functions",
        "16. Prediction, batch scoring, and model deployment",
        "17. API and implementation function reference",
        "18. Security, testing, and deployment boundaries",
        "19. Installation, operation, and maintenance",
        "Appendix B. Project files and runtime responsibilities",
    ]
    for entry in entries:
        page_break.insert_paragraph_before(entry, style="Normal")


def append_extension(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    add_heading(doc, "Part Two CreditScope Web Application", 1)
    add_body(
        doc,
        "The notebook work was extended into CreditScope, a local research application that connects the saved Home Credit analysis to an interactive dashboard, deployable model registry, individual and batch prediction workflows, downloadable reports, and optional AI assistance. The application keeps the notebook evidence visible while separating historical observations from live model scores.",
    )

    add_heading(doc, "13 Productization from Notebook to Application", 1)
    add_heading(doc, "Purpose of the application extension", 2)
    add_body(
        doc,
        "The notebooks establish the data preparation, exploratory analysis, engineered features, model comparisons, and final submission logic. Productization adds the runtime pieces needed to browse those results, load trained artifacts safely, rebuild model inputs in a fixed order, score new cases, process CSV batches, and expose the same capabilities through a browser and API.",
    )
    add_bullets(doc, [
        "Keep the full 307,511-row applicant snapshot available for interactive cohort analysis.",
        "Preserve notebook-derived research tables, plots, preprocessing stages, feature lineage, and model metrics as read-only application assets.",
        "Allow multiple trained model formats to be deployed only after schema and inference validation.",
        "Show assumptions and local sensitivities with every applicant score so users can see which values were entered, derived, supplied by CSV, or filled from training defaults.",
        "Treat CreditScope as a local research tool rather than an automated credit-decision system.",
    ])

    add_heading(doc, "Repository data and generated assets", 2)
    add_table(doc, ["Asset", "Created or loaded by", "Application use"], [
        ["data/applicants.pkl", "scripts/import_project.py", "Full applicant snapshot used by dashboard filters and peer-cohort calculations"],
        ["data/schema.json", "Import/export process", "Ordered feature defaults, methods, quantile boundaries, and model input schema"],
        ["data/provenance.json", "Import process", "Source and preparation notes displayed by the application"],
        ["data/gallery.json and data/eda", "Notebook plot import", "Catalog and files for all 45 saved EDA/model figures"],
        ["data/research/*.csv", "Notebook extraction", "Feature lineage, preprocessing stages, correlations, model metrics, curves, thresholds, and ensemble results"],
        ["data/global_risk.json", "Project data asset", "Country-level banking and macroeconomic context used by Worldwide statistics"],
        ["storage/models", "Deploy model workflow", "Validated local model packages available for prediction"],
        ["storage/archived", "Model removal workflow", "Recoverable archive for models removed from the active registry"],
    ], [1.55, 1.65, 3.9])

    add_heading(doc, "Notebook to runtime sequence", 2)
    add_numbers(doc, [
        "Run the notebook pipeline and save the trained artifact, ordered feature list, preprocessing object when required, threshold, and validation metadata.",
        "Import the notebook and Power BI outputs into stable local assets without executing notebook code inside the web server.",
        "Start the FastAPI application and load the applicant snapshot, schema, research tables, gallery index, and model registry on demand.",
        "Select a deployed model and construct its exact ordered feature row from a form submission or CSV row.",
        "Apply the matching preprocessor when the model package stores preprocessing separately, then calculate the class-1 score.",
        "Compare the score with the saved validation threshold and return the flag, assumptions, local sensitivities, and model caveats.",
        "Present results in the browser, export JSON or CSV, or send verified case context to the selected assistant.",
    ])

    add_heading(doc, "14 CreditScope Architecture and Execution Flow", 1)
    add_heading(doc, "Application layers", 2)
    add_table(doc, ["Layer", "Main components", "Responsibility"], [
        ["Browser interface", "frontend/index.html, styles.css, app.js", "Navigation, filters, charts, forms, uploads, result rendering, downloads, and assistant interaction"],
        ["Static documentation", "frontend/docs.html and docs.js", "Local interactive API reference at /docs"],
        ["HTTP application", "backend/main.py", "FastAPI routes, request validation, upload handling, security middleware, assistant connections, and static-file serving"],
        ["Analytics", "backend/analytics.py", "Cached applicant data, six filters, cohort summaries, grouped rates, and research CSV access"],
        ["Feature construction", "backend/features.py", "Twelve-field applicant schema, derived variables, defaults, categorical handling, and exact feature ordering"],
        ["Model runtime", "backend/models.py and backend/dcn.py", "Registry, artifact loading, preprocessing, multi-framework inference, thresholding, and sensitivity calculations"],
        ["Batch runtime", "backend/batch.py", "CSV parsing, row validation, model-column pass-through, temporary session storage, and selected-row retrieval"],
        ["Report export", "backend/reporting.py", "Two-page PDF dashboard summary built from the active filter state"],
        ["Local data and models", "data and storage directories", "Imported analytical assets, active model packages, and recoverable model archives"],
    ], [1.35, 2.2, 3.55])

    add_heading(doc, "Request and response flow", 2)
    add_numbers(doc, [
        "The browser calls the shared api function, which prefixes /api, adds JSON headers or an administrator token when needed, and surfaces server errors as user messages.",
        "FastAPI validates host, origin, request schema, file type, size, field ranges, and model identifier before application logic runs.",
        "The selected backend module calculates analytics, builds feature rows, loads a cached model, scores applicants, or prepares assistant context.",
        "The API returns JSON for interactive screens, CSV/JSON content for client-side downloads, or a streamed PDF for dashboard reporting.",
        "Frontend render functions update the active page and bind buttons, dialogs, filters, charts, downloads, and navigation events.",
    ])

    add_heading(doc, "Runtime stack", 2)
    add_table(doc, ["Technology", "Use in CreditScope"], [
        ["FastAPI and Pydantic", "API routing, validation, file upload forms, and generated OpenAPI schema"],
        ["pandas and NumPy", "Applicant analytics, feature rows, CSV batch preparation, and model inputs"],
        ["scikit-learn and joblib", "Pipeline/estimator loading and fitted preprocessing"],
        ["LightGBM and CatBoost", "Native loading of text and CBM model artifacts"],
        ["PyTorch", "DCN checkpoint and TorchScript inference on CPU"],
        ["TensorFlow optional runtime", "Safe loading of complete Keras models when requirements-keras.txt is installed"],
        ["ReportLab", "Filtered executive dashboard PDF export"],
        ["HTML, CSS, and vanilla JavaScript", "Responsive single-page browser application without a frontend framework"],
        ["Ollama and Google Gemini", "Optional local or cloud assistant backends"],
    ], [2.05, 5.05])

    add_heading(doc, "15 Website Pages and User Functions", 1)
    add_body(
        doc,
        "CreditScope uses hash-based navigation and keeps shared application state in the browser. Each page requests only the data it needs and reuses the same loaded configuration, model registry, dashboard filters, selected case, and assistant histories.",
    )
    add_heading(doc, "Seven navigation sections", 2)
    add_table(doc, ["Page", "Functions available to the user"], [
        ["Executive overview", "View the full portfolio or filter cohorts; review counts, rates, averages, exposure, age/debt/income/education profiles, age-income matrix, scenario loss, and export a PDF summary"],
        ["Worldwide statistics", "Select a country on an interactive Robinson-projection map; compare World Bank indicators and a sourced banking exposure scenario"],
        ["Individual prediction", "Enter 12 applicant fields, choose a deployed model, obtain a score and threshold flag, inspect sensitivities and full feature provenance, compare models, and export JSON"],
        ["Batch CSV check", "Upload a named-column or 107-feature-order CSV, score as many as 50,000 rows, review errors and ranked results, export results, and inspect one row in detail"],
        ["Explore the data", "Browse 45 saved plots, filter the gallery, review preprocessing stages, view the 11-model leaderboard, and inspect research metrics"],
        ["AI assistant", "Use Ollama or Gemini for project questions and optionally attach a verified scored applicant plus historical peer cohort"],
        ["Deploy model", "Upload trusted model packages, validate metadata and feature order, run a reference-case smoke test, activate successful models, and archive removed models"],
    ], [1.55, 5.55])

    add_heading(doc, "Executive overview", 2)
    add_body(
        doc,
        "The dashboard starts with all 307,511 applicant records. Six filters are available: age group, gender, education, income type, bureau overdue status, and debt-to-income group. The backend applies selected values to the snapshot and recalculates every metric, group chart, and matrix from the resulting cohort.",
    )
    add_bullets(doc, [
        "Headline measures: applicant count, recorded repayment-difficulty cases and rate, average income, average credit, total requested-credit exposure, difficulty exposure, average age, and average credit-to-income ratio.",
        "Breakdowns: age group, debt-to-income group, income type, education, and a two-dimensional age by income-quintile matrix.",
        "Scenario calculation: users can combine the observed rate or another displayed score with loss given default and exposure assumptions. The result is a scenario estimate, not an accounting provision.",
        "Export: /api/dashboard/report creates a two-page PDF using the same active filters and excludes applicant-level rows.",
    ])

    add_heading(doc, "Worldwide statistics", 2)
    add_body(
        doc,
        "The worldwide page combines a local map asset with a stored country snapshot. It displays bank non-performing loans, GDP growth, inflation, unemployment, lending interest rate, bank capital to assets, and private-sector credit relative to GDP. Metric information dialogs state what each value means and how it should be interpreted. External links identify the World Bank indicators and the Natural Earth map source.",
    )

    add_heading(doc, "Explore the data", 2)
    add_body(
        doc,
        "The research page loads the plot catalog from gallery.json and the notebook-derived insight, preprocessing, and model-metric tables from data/research. Users can search and filter plots, switch among research groupings, read preprocessing stages, and compare the 11 recorded models without rerunning the notebooks.",
    )

    add_heading(doc, "AI assistant", 2)
    add_body(
        doc,
        "The assistant page supports a local Ollama model and Google Gemini. Each assistant keeps a separate browser-memory conversation. General questions receive aggregate dashboard, notebook, research-metric, EDA, and model-registry context. When a case is attached, the server adds its verified model score, threshold, entered values, top local sensitivities, the portfolio observed rate, and a peer cohort defined by age within five years and income within 25 percent.",
    )
    add_bullets(doc, [
        "Ollama defaults to gemma3:4b at a loopback URL and is rejected if configured to a non-local hostname.",
        "Gemini uses GEMINI_API_KEY and the configured Gemini model; the key remains server-side.",
        "Assistant instructions require evidence-grounded answers and prohibit inventing missing model files or results.",
        "Selected case context can be cleared before continuing a general conversation.",
    ])

    add_heading(doc, "16 Prediction Batch Scoring and Model Deployment", 1)
    add_heading(doc, "Individual applicant inputs", 2)
    add_table(doc, ["Input", "Validation or transformation"], [
        ["Age", "19 to 100; used directly and converted to the notebook age band"],
        ["Annual income", "1,000 to 100,000,000; used in credit/income and debt/income derivatives"],
        ["Credit amount", "1,000 to 100,000,000; used directly and converted to the notebook credit quantile"],
        ["Loan annuity", "1 to 10,000,000"],
        ["Employment duration", "0 to 85 years and cannot imply employment before age 14; converted to negative DAYS_EMPLOYED"],
        ["Education", "Five ordered levels mapped to the notebook ordinal code"],
        ["Income type", "Working, Commercial associate, Pensioner, State servant, or Other; rebuilds one-hot fields where required"],
        ["Contract type", "Cash loans or Revolving loans mapped to 0 or 1"],
        ["Car ownership", "Boolean mapped to FLAG_OWN_CAR"],
        ["External score 2", "Optional value from 0 to 1; otherwise the training default is used"],
        ["External score 3", "Optional value from 0 to 1; otherwise the training default is used"],
        ["Late payment rate", "Optional value from 0 to 1; otherwise the training default is used"],
    ], [2.05, 5.05])

    add_heading(doc, "Exact online feature reconstruction", 2)
    add_body(
        doc,
        "make_model_features builds the exact ordered input declared by the selected model. Entered values override defaults, derived groups and ratios are recomputed, supplied CSV model columns can pass through, native CatBoost categoricals remain strings, and unsupported feature names stop the request. Numeric missing history is filled from schema defaults; low-cardinality categoricals use training modes when the form cannot provide them.",
    )
    add_numbers(doc, [
        "Validate the applicant with the Pydantic Applicant model.",
        "Read the model metadata and feature order from its active registry folder.",
        "Build one numeric or mixed categorical row in exactly that order and record the source of every value.",
        "Load and cache the model and fitted preprocessor required by its framework.",
        "Calculate the class-1 score and verify that it is finite and between 0 and 1.",
        "Compare the score with the saved threshold and return the research flag.",
        "Replace one entered input or supplied feature at a time with its reference value to calculate local sensitivity differences.",
    ])
    add_body(
        doc,
        "The sensitivity values are one-at-a-time counterfactual score differences. They are not SHAP values, additive contributions, causal explanations, or evidence that changing one field would change an actual lending outcome. Scores from weighted or SMOTE-trained models are also not assumed to be calibrated population default probabilities.",
    )

    add_heading(doc, "Batch CSV workflow", 2)
    add_bullets(doc, [
        "Accepted files must be UTF-8 CSV data and no larger than 60 MB.",
        "The parser detects comma, tab, semicolon, or pipe delimiters and rejects inconsistent row widths or duplicate/blank headers.",
        "Named-column files may use human-readable input names, raw Home Credit names, or exact model-feature names.",
        "Headerless numeric files are accepted only when they contain exactly the 107 columns in CreditScope feature order.",
        "A batch can contain at most 50,000 applicants. Invalid rows are reported while valid rows continue to scoring.",
        "Results include row number, applicant identifier, exposure, score, threshold flag, and assumed-feature count.",
        "The server keeps at most two batch sessions in memory for one hour so a user can open one selected row with full feature provenance and sensitivities.",
    ])

    add_heading(doc, "Supported deployment formats", 2)
    add_table(doc, ["Model kind", "Artifact and required companion files", "Validation performed"], [
        ["Scikit-learn", ".joblib or .pkl; metadata; ordered features; fitted preprocessor unless the pipeline contains it or metadata declares none", "predict_proba, classes [0,1], feature count/order, and reference prediction"],
        ["PyTorch DCN checkpoint", ".pth or .pt checkpoint plus matching preprocessor", "Checkpoint feature order, DCN state dictionary, input dimension, threshold, and reference prediction"],
        ["PyTorch TorchScript", ".pt plus matching preprocessor and metadata", "TorchScript load, input transformation, output range, and reference prediction"],
        ["LightGBM", ".txt plus metadata and ordered features", "Native Booster load, feature count, output range, and reference prediction"],
        ["CatBoost", ".cbm plus metadata, ordered features, and categorical feature list", "Native model load, feature-name order, categorical membership, output range, and reference prediction"],
        ["Keras", "Complete .keras or .h5 model; metadata; ordered features; optional fitted preprocessor", "Safe-mode model load, input width, output shape/range, and reference prediction"],
    ], [1.35, 3.35, 2.4], 8.2)

    add_heading(doc, "Deployment and removal behavior", 2)
    add_body(
        doc,
        "A deployment request must identify the model kind, include a trusted-files confirmation, and provide compatible metadata. Uploaded files are first written to a temporary directory. CreditScope validates file extensions, feature names, threshold, preprocessors, categorical declarations, model structure, and a real reference-case inference. Only a successful package is moved into storage/models and exposed to prediction. Removing a model clears its cache and moves its folder to storage/archived; original notebook artifacts outside the application are never deleted.",
    )

    add_heading(doc, "17 API and Implementation Function Reference", 1)
    add_heading(doc, "HTTP endpoints", 2)
    add_table(doc, ["Method and path", "Purpose"], [
        ["GET /api/health", "Return data readiness, assistant configuration, selected model names, and administrator-token status"],
        ["GET /api/config", "Return dashboard filter options, reference applicant, provenance, global indicators, and health information"],
        ["GET /api/dashboard", "Apply query-string cohort filters and return recalculated dashboard measures and groups"],
        ["GET /api/dashboard/report", "Create a filtered two-page PDF executive dashboard report"],
        ["GET /api/eda", "Return plot catalog, notebook insights, preprocessing stages, and model metrics"],
        ["GET /api/models", "List active validated model packages"],
        ["GET /api/models/template", "Return starter model metadata for a CreditScope numeric model"],
        ["POST /api/models", "Validate and deploy a trusted model package"],
        ["DELETE /api/models/{model_id}", "Archive an active model and remove it from prediction"],
        ["POST /api/predict", "Score one validated applicant with one deployed model"],
        ["POST /api/compare", "Score one applicant across every compatible active model"],
        ["POST /api/batch/predict", "Parse, validate, and score a CSV batch"],
        ["POST /api/batch/detail", "Reconstruct one retained batch row and return full prediction details"],
        ["POST /api/chat", "Send project or selected-applicant context to Ollama or Gemini"],
        ["GET /", "Serve the CreditScope browser application"],
        ["GET /docs", "Serve the local API reference interface"],
    ], [2.3, 4.8])

    add_heading(doc, "Backend functions", 2)
    backend_rows = [
        ["analytics.py", "applicants", "Load and cache the applicant snapshot"],
        ["analytics.py", "records", "Convert a DataFrame to JSON-ready records"],
        ["analytics.py", "research", "Load a named research CSV"],
        ["analytics.py", "group", "Calculate count, positive cases, and rate by one column"],
        ["analytics.py", "dashboard", "Apply six filters and compute all dashboard measures, groups, and matrix cells"],
        ["analytics.py", "filter_options", "Return distinct values for each dashboard filter"],
        ["features.py", "Applicant", "Validate the 12 human-readable input fields"],
        ["features.py", "schema", "Load saved defaults, methods, bins, and features"],
        ["features.py", "bucket", "Apply right-closed notebook-style bins and clamp out-of-range inputs"],
        ["features.py", "make_features", "Build supported numeric features and provenance from one applicant"],
        ["features.py", "category_mode", "Return a cached training mode for a categorical field"],
        ["features.py", "category_frequency", "Return cached training frequencies for high-cardinality categories"],
        ["features.py", "make_model_features", "Build the exact ordered numeric/native-categorical model row"],
        ["features.py", "reference_case", "Construct the baseline applicant from saved training defaults"],
        ["models.py", "read_metadata", "Read one deployed model's metadata"],
        ["models.py", "registry", "List valid active model folders in name order"],
        ["models.py", "model_folder", "Validate a model identifier and locate its active folder"],
        ["models.py", "load", "Load, validate, smoke-test, and cache a framework-specific predictor"],
        ["models.py", "transform", "Apply a saved pipeline, scaler, or imputer/scaler pair"],
        ["models.py", "predict_case", "Score one case, apply the threshold, expose assumptions, and calculate local sensitivities"],
        ["batch.py", "remember", "Store up to two recent batches and remove sessions older than one hour"],
        ["batch.py", "selected", "Retrieve one row from an unexpired batch session"],
        ["batch.py", "parse_csv", "Detect delimiter/header mode and validate batch dimensions and encoding"],
        ["batch.py", "recognized_columns", "Find uploaded columns that can affect the applicant or selected model"],
        ["batch.py", "applicant_from_row", "Map human-readable, raw, or encoded CSV fields into Applicant"],
        ["batch.py", "model_values", "Pass through exact model columns and encode selected raw categoricals"],
        ["dcn.py", "CrossLayer", "Implement one explicit cross-feature layer"],
        ["dcn.py", "DeepCrossNetwork", "Combine three cross layers with the deep network used by the checkpoint adapter"],
        ["reporting.py", "build_dashboard_report", "Build the filtered dashboard PDF in memory"],
        ["main.py", "boundaries", "Enforce local host/origin boundaries and response security headers"],
        ["main.py", "admin", "Require the optional administrator token for protected operations"],
        ["main.py", "health and config", "Describe readiness, assistants, filters, provenance, reference case, and global data"],
        ["main.py", "dashboard and dashboard_report", "Serve dashboard JSON and the filtered PDF"],
        ["main.py", "eda and model_list", "Serve research assets and the current model registry"],
        ["main.py", "batch_predict and batch_detail", "Score uploaded CSV rows and inspect a retained row"],
        ["main.py", "predict and compare", "Score one model or compare all compatible models"],
        ["main.py", "save_upload and deploy", "Limit, validate, smoke-test, and activate uploaded artifacts"],
        ["main.py", "delete", "Move an active model into the recoverable archive"],
        ["main.py", "chat_instructions and applicant_brief", "Create evidence-grounded assistant context"],
        ["main.py", "ollama_url and chat", "Restrict the local provider URL and call Ollama or Gemini"],
        ["main.py", "index and api_docs", "Serve the main application and local API documentation"],
    ]
    add_table(doc, ["Module", "Function or class", "Responsibility"], backend_rows, [1.25, 2.1, 3.75], 7.8)

    add_heading(doc, "Frontend functions", 2)
    frontend_rows = [
        ["api", "Send API requests, attach JSON/admin headers, parse responses, and raise readable errors"],
        ["download and toast", "Create local downloads and display temporary status messages"],
        ["heading, stat, panel, field", "Generate reusable page headings, metric cards, panels, and form controls"],
        ["infoButton and bindInfoButtons", "Open plain-language metric definitions and interpretation guidance"],
        ["bars and matrix", "Render grouped rates and the age-income matrix"],
        ["robinsonPoint", "Project stored latitude/longitude coordinates onto the browser map"],
        ["indicatorCard and globalRiskMarkup", "Render country indicators, source notes, and the world map"],
        ["defaultRiskPulseMarkup", "Show the selected country's non-performing-loan position"],
        ["lossPanelMarkup and bindLossPanel", "Render and update expected-loss scenario controls"],
        ["portfolioLossMarkup and bindGlobalRisk", "Apply scenario logic to portfolio and global pages"],
        ["renderGlobal", "Load and display Worldwide statistics"],
        ["renderOverview and refreshDashboard", "Render the executive overview and recalculate after filters"],
        ["readCase", "Read the 12 applicant form values"],
        ["renderPredict", "Build the prediction screen and submit or compare a case"],
        ["selectChatCase", "Attach a scored case to the assistant workflow"],
        ["renderBatch and renderBatchResults", "Upload CSV data, display ranked results/errors, and export results"],
        ["resultMarkup and bindResult", "Display score, threshold, sensitivities, assumptions, exports, and follow-up actions"],
        ["renderEDA and renderEDABody", "Load and filter the notebook research page"],
        ["researchMetrics and renderGallery", "Summarize metrics and render plot cards"],
        ["currentAssistant and formattedAnswer", "Select an assistant and safely convert its Markdown response"],
        ["renderCaseChat, renderAssistant, renderMessages", "Render assistant selection, selected-case context, and chat history"],
        ["sendChat", "Submit a general or case-grounded assistant request"],
        ["adminInput and bindAdmin", "Keep the optional administrator token in page memory"],
        ["renderModels and refreshModels", "Render model deployment/removal controls and refresh the active registry"],
        ["currentPage, navigate, start", "Control hash navigation, route rendering, initial configuration, and startup"],
    ]
    add_table(doc, ["Function group", "Responsibility"], frontend_rows, [2.35, 4.75], 8.0)

    add_heading(doc, "18 Security Testing and Deployment Boundaries", 1)
    add_heading(doc, "Implemented local controls", 2)
    add_bullets(doc, [
        "Requests are limited to localhost, loopback addresses, and the test host unless a secured reverse proxy is added.",
        "Cross-origin POST, PUT, PATCH, and DELETE requests are rejected.",
        "Responses include content-type sniffing, referrer, framing, and content-security-policy headers.",
        "Pydantic forbids unexpected applicant fields, rejects non-finite numbers, and enforces numeric ranges and employment-age consistency.",
        "Upload size, extension, metadata, feature order, framework requirements, threshold, and actual inference output are checked before activation.",
        "Serialized joblib and pickle files are accepted only after the user confirms they are trusted because loading them can execute code.",
        "An optional CREDITSCOPE_ADMIN_TOKEN protects model mutation and AI use; remote model management is blocked when no token is configured.",
        "Removed models are archived instead of permanently deleted.",
    ])

    add_heading(doc, "Automated verification", 2)
    add_body(
        doc,
        "The current repository test run completed 16 tests successfully. The suite isolates the model registry and checks dashboard totals and empty cohorts, feature transformations and ordering, real MLP parity, original DCN loading, trusted upload and archival, invalid inputs, host and origin protections, assistant setup errors, batch behavior, selected-applicant context, and chart assets. Artifact-dependent checks are skipped only when the corresponding original model files are unavailable.",
    )
    add_code_block(doc, ".\\.venv\\Scripts\\python.exe -m pytest tests -q")

    add_heading(doc, "Boundaries before remote hosting", 2)
    add_table(doc, ["Current local design", "Required for an internet-facing service"], [
        ["Single local user with optional administrator token", "Full authentication, authorization, roles, session controls, and audit logging"],
        ["Loopback host and same-origin mutation checks", "HTTPS, secured reverse proxy, explicit trusted hosts/origins, and network controls"],
        ["In-process model loading", "Isolated model-loading workers or containers with resource and filesystem restrictions"],
        ["Basic request and upload limits", "Rate limiting, quotas, malware scanning, timeout policy, and abuse monitoring"],
        ["Local files for models and archives", "Backups, retention policy, encryption, disaster recovery, and controlled artifact promotion"],
        ["Research applicant snapshot", "Reviewed borrower-data governance, access policy, minimization, consent/legal basis, and deletion procedures"],
        ["Local operational visibility", "Central logs, metrics, alerts, model drift monitoring, incident response, and service-level objectives"],
    ], [3.2, 3.9], 8.3)

    add_heading(doc, "19 Installation Operation and Maintenance", 1)
    add_heading(doc, "Clone and start on Windows", 2)
    add_code_block(doc, "git lfs install\ngit clone https://github.com/MImahin/CreditScope.git\ncd CreditScope\ngit lfs pull\npython -m venv .venv\n.\\.venv\\Scripts\\python.exe -m pip install --upgrade pip\n.\\.venv\\Scripts\\python.exe -m pip install -r requirements.txt\nCopy-Item .env.example .env\n.\\start.bat")
    add_body(
        doc,
        "After the launcher reports that startup is complete, open http://127.0.0.1:8000 and keep the launcher window running. The launcher checks the environment and imports the project data if the local snapshot is missing. Press Ctrl+C in the launcher window to stop the service.",
    )

    add_heading(doc, "Optional configuration", 2)
    add_table(doc, ["Setting", "Purpose"], [
        ["CREDITSCOPE_SOURCE", "Use a different original dataset folder when rebuilding imported assets"],
        ["CREDITSCOPE_ADMIN_TOKEN", "Protect model mutation and assistant requests"],
        ["CHAT_PROVIDER", "Select the default assistant provider"],
        ["OLLAMA_BASE_URL", "Local Ollama base URL; only loopback hostnames are accepted"],
        ["OLLAMA_MODEL", "Local assistant model name, default gemma3:4b"],
        ["GEMINI_API_KEY", "Enable the Gemini assistant"],
        ["GEMINI_MODEL", "Gemini model name, default gemini-3.5-flash-lite"],
    ], [2.25, 4.85])

    add_heading(doc, "Routine operating checklist", 2)
    add_numbers(doc, [
        "Confirm data/schema.json and data/applicants.pkl are present, or rebuild them with scripts/import_project.py.",
        "Start CreditScope and verify /api/health reports status ok and data_ready true.",
        "Deploy each trained model with the exact artifact, feature list, metadata, categorical list, and preprocessor from one package folder.",
        "Run an individual reference case and a small CSV batch before a demonstration or analysis session.",
        "Use the dashboard source notes and model caveats when presenting results; do not describe a model score as a calibrated probability unless its metadata supports that claim.",
        "Run the automated tests after modifying feature construction, model loading, API validation, security middleware, or frontend assets.",
        "Archive obsolete models through the application and back up storage/models, storage/archived, and any updated data assets before moving computers.",
    ])

    add_heading(doc, "Rebuilding imported research data", 2)
    add_code_block(doc, ".\\.venv\\Scripts\\python.exe scripts\\import_project.py")
    add_body(
        doc,
        "The importer reads the original project sources and writes the local application snapshot. It does not execute notebook code or modify the source files. The dashboard is therefore based on a local snapshot rather than a live Power BI connection; rerun the importer and restart the application when source data changes. Original amount units are preserved and no currency conversion is inferred.",
    )

    add_heading(doc, "Appendix B Project Files and Runtime Responsibilities", 1)
    add_table(doc, ["Path", "Role"], [
        ["backend/main.py", "FastAPI application, routes, middleware, model upload/removal, chat providers, and static serving"],
        ["backend/analytics.py", "Applicant snapshot loading, filters, cohort metrics, grouped rates, and research table access"],
        ["backend/features.py", "Applicant validation, defaults, feature derivation, categoricals, aliases, and ordered model rows"],
        ["backend/models.py", "Active registry, multi-framework loaders, preprocessors, scoring, threshold flags, and sensitivities"],
        ["backend/batch.py", "CSV parsing, row mapping, model-value pass-through, and short-lived batch sessions"],
        ["backend/dcn.py", "Deep and Cross Network architecture for original checkpoint compatibility"],
        ["backend/reporting.py", "Filtered executive dashboard PDF generation"],
        ["frontend/index.html", "Application shell, dialogs, navigation, and page mount points"],
        ["frontend/styles.css", "Responsive visual system for the dashboard and forms"],
        ["frontend/app.js", "State, page rendering, interactions, API calls, downloads, maps, charts, prediction, batch, registry, and chat"],
        ["frontend/docs.html and docs.js", "Local OpenAPI documentation interface"],
        ["scripts/import_project.py", "Build application data and research assets from the original project sources"],
        ["scripts/export_smote_models.py", "Package already-trained SMOTE estimators for deployment without retraining or inference-time resampling"],
        ["tests/test_api.py", "Integration tests using temporary model registries"],
        ["start.bat", "Windows launcher and environment/data readiness checks"],
        ["requirements.txt", "Core Python runtime dependencies"],
        ["requirements-keras.txt", "Optional TensorFlow dependency for Keras models"],
        ["README.md", "User setup, model deployment, interpretation, testing, and boundary guidance"],
    ], [2.45, 4.65], 8.2)

    add_heading(doc, "End to end completion checklist", 2)
    add_numbers(doc, [
        "Complete the notebook pipeline and retain its feature, metric, and caveat evidence.",
        "Export each production model with its exact ordered input schema, preprocessing object, threshold, and validation metadata.",
        "Import the applicant snapshot, plots, research tables, provenance, and global context into the CreditScope data directory.",
        "Start the local FastAPI service and confirm dashboard and EDA assets load.",
        "Deploy trusted model packages and confirm the reference inference smoke test passes.",
        "Verify individual scoring, comparison, feature provenance, local sensitivities, batch upload/detail, and downloads.",
        "Verify Ollama and Gemini separately if those optional assistants will be demonstrated.",
        "Run the test suite and review the security and research-use boundaries before sharing the application.",
    ])


def main():
    parser = argparse.ArgumentParser(description="Extend a notebook-grounded Word report with CreditScope application documentation.")
    parser.add_argument("source",nargs="?",type=Path,default=DEFAULT_SOURCE,
                        help="Source notebook report (.docx). Defaults to a file beside the project README.")
    parser.add_argument("--output",type=Path,default=DEFAULT_OUTPUT,
                        help="Destination .docx path.")
    args = parser.parse_args()
    source=args.source.expanduser().resolve()
    output=args.output.expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    output.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(source,output)
    doc = Document(output)
    update_front_matter(doc)
    append_extension(doc)
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.core_properties.title = "Home Credit Data Analytics and CreditScope Web Application Technical Report"
    doc.core_properties.subject = "Notebook-grounded analysis with the implemented CreditScope web application"
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
